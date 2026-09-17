import os
from datetime import datetime
import openpyxl
from konfigurasi import EXCEL_FILE_REJECT as DEFAULT_EXCEL_FILE


class ExcelManager:
    """Mengelola pembacaan data reject dan penulisan status kembali ke file Excel."""

    def __init__(self, file_path=None):
        self.file_path = file_path or DEFAULT_EXCEL_FILE

    def baca_data_reject(self) -> list:
        """
        Membaca data IDPEL, NOMETER, NIK, NAMA, dan NO TELP dari file Excel.
        Melewati (skip) baris yang statusnya sudah selesai / sukses.
        """
        if not os.path.exists(self.file_path):
            print(f"[ERROR] File Excel '{self.file_path}' tidak ditemukan.")
            return []

        print(f"[EXCEL] Membaca data dari '{self.file_path}'...")
        data_list = []
        skipped_count = 0
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]
            idpel_idx = -1
            nometer_idx = -1
            nik_idx = -1
            nama_idx = -1
            no_telp_idx = -1
            status_idx = -1

            for idx, h in enumerate(headers):
                if h is None:
                    continue
                h_str = str(h).strip().upper()
                if h_str == "IDPEL":
                    idpel_idx = idx
                elif h_str in ["NOMETER_BARU", "NOMETER", "NO METER", "NO_METER"]:
                    nometer_idx = idx
                elif h_str == "NIK":
                    nik_idx = idx
                elif h_str == "NAMA":
                    nama_idx = idx
                elif h_str in ["NO TELP", "NO TELP/HP", "TELP", "TELEPON"]:
                    no_telp_idx = idx
                elif h_str == "STATUS":
                    status_idx = idx

            if idpel_idx == -1:
                print("[ERROR] Kolom 'IDPEL' tidak ditemukan di header Excel.")
                wb.close()
                return []

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                idpel_val = row[idpel_idx] if idpel_idx < len(row) else None
                nometer_val = row[nometer_idx] if (nometer_idx != -1 and nometer_idx < len(row)) else None
                nik_val = row[nik_idx] if (nik_idx != -1 and nik_idx < len(row)) else None
                nama_val = row[nama_idx] if (nama_idx != -1 and nama_idx < len(row)) else None
                no_telp_val = row[no_telp_idx] if (no_telp_idx != -1 and no_telp_idx < len(row)) else None
                status_val = row[status_idx] if (status_idx != -1 and status_idx < len(row)) else None

                # Skip baris yang statusnya sudah memuat 'SUKSES', 'SUBMIT', 'TIDAK DITEMUKAN', 'ERROR', atau 'KOORDINAT'
                if status_val is not None:
                    status_str = str(status_val).strip().upper()
                    if "ALAMAT TIDAK DITEMUKAN" not in status_str and "ALAMAT NULL" not in status_str:
                        if any(x in status_str for x in ["SUKSES", "SUBMIT", "TIDAK DITEMUKAN", "ERROR", "KOORDINAT & FOTO TIDAK ADA", "KOORDINAT"]):
                            skipped_count += 1
                            continue

                if idpel_val is not None:
                    idpel_str = str(idpel_val).strip()
                    if idpel_str and idpel_str.lower() != "none":
                        nometer_str = str(nometer_val).strip() if nometer_val is not None else ""
                        nik_str = str(nik_val).strip() if nik_val is not None else ""
                        nama_str = str(nama_val).strip() if nama_val is not None else ""
                        no_telp_str = str(no_telp_val).strip() if no_telp_val is not None else "-"
                        data_list.append({
                            "row": row_idx,
                            "idpel": idpel_str,
                            "nometer": nometer_str,
                            "nik": nik_str,
                            "nama": nama_str,
                            "no_telp": no_telp_str
                        })

            wb.close()
            print(f"[EXCEL] Berhasil membaca {len(data_list)} entri data yang perlu diproses ({skipped_count} data berstatus sukses/selesai dilewati).")
        except Exception as e:
            print(f"[ERROR] Gagal membaca file Excel: {e}")

        return data_list

    def simpan_status(self, row_num: int, status_text: str = "SUKSES DIUPDATE") -> bool:
        """
        Menyimpan status pengerjaan data ke kolom 'STATUS' dan 'TGL UPDATE STATUS' pada file Excel.
        """
        if not os.path.exists(self.file_path):
            print(f"[ERROR] File Excel '{self.file_path}' tidak ditemukan untuk simpan status.")
            return False

        try:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb.active

            status_col = -1
            tgl_col = -1

            for col_idx, cell in enumerate(ws[1], start=1):
                if cell.value is not None:
                    cell_val = str(cell.value).strip().upper()
                    if cell_val == "STATUS":
                        status_col = col_idx
                    elif "TGL" in cell_val or "TANGGAL" in cell_val or "UPDATE" in cell_val:
                        tgl_col = col_idx

            if status_col == -1:
                status_col = ws.max_column + 1
                ws.cell(row=1, column=status_col, value="STATUS")
                print(f"[EXCEL] Membuat kolom baru 'STATUS' pada kolom ke-{status_col}")

            ws.cell(row=row_num, column=status_col, value=status_text)

            waktu_sekarang = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            if tgl_col != -1:
                ws.cell(row=row_num, column=tgl_col, value=waktu_sekarang)
                print(f"[EXCEL] Berhasil menyimpan status '{status_text}' & tanggal '{waktu_sekarang}' pada baris {row_num} di Excel.")
            else:
                print(f"[EXCEL] Berhasil menyimpan status '{status_text}' pada baris {row_num} di Excel.")

            wb.save(self.file_path)
            wb.close()
            return True
        except PermissionError:
            print(f"[ERROR EXCEL] File '{self.file_path}' sedang dibuka di Microsoft Excel! Harap TUTUP file Excel agar status dapat tersimpan.")
            return False
        except Exception as err:
            print(f"[ERROR EXCEL] Gagal menyimpan status ke Excel pada baris {row_num}: {err}")
            return False
